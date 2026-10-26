from __future__ import annotations

from copy import copy, deepcopy
from hashlib import sha256
from pathlib import Path
import re
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(r"C:\Users\29580\Desktop\就业\最新简历\2026张宇涵求职简历AI应用开发.docx")
OUTPUT = ROOT / "output" / "docx" / "2026张宇涵求职简历AI应用开发_新增工业图纸项目.docx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
NS = {"w": W}


def tag(local):
    return f"{{{W}}}{local}"


def text_of(element):
    return "".join(element.itertext()) if element is None else "".join(element.xpath(".//w:t/text()", namespaces=NS))


def set_val(parent, name, value="1"):
    element = parent.find(tag(name))
    if element is None:
        element = etree.SubElement(parent, tag(name))
    element.set(tag("val"), str(value))
    return element


def new_paragraph(prototype, text, *, kind="body", keep_next=True):
    paragraph = etree.Element(tag("p"))
    paragraph.set(f"{{{W14}}}paraId", uuid4().hex[:8].upper())
    ppr = deepcopy(prototype.find(tag("pPr")))
    if ppr is None:
        ppr = etree.Element(tag("pPr"))
    paragraph.append(ppr)
    if kind != "bullet":
        for element in list(ppr):
            if etree.QName(element).localname in {"numPr", "ind"}:
                ppr.remove(element)
        indent = etree.SubElement(ppr, tag("ind"))
        indent.set(tag("left"), "0")
        indent.set(tag("firstLine"), "0")
        indent.set(tag("firstLineChars"), "0")
    set_val(ppr, "jc", "left")
    set_val(ppr, "keepNext", "0")
    set_val(ppr, "keepLines", "1")
    set_val(ppr, "widowControl", "1")
    set_val(ppr, "snapToGrid", "0")
    spacing = ppr.find(tag("spacing"))
    if spacing is None:
        spacing = etree.SubElement(ppr, tag("spacing"))
    line = 310 if kind in {"title", "bullet", "body"} else 270
    spacing.attrib.update({tag("before"): "0", tag("after"): "30", tag("line"): str(line), tag("lineRule"): "exact"})
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        bold = kind == "title" or (part.startswith("**") and part.endswith("**"))
        content = part[2:-2] if part.startswith("**") and part.endswith("**") else part
        run = etree.SubElement(paragraph, tag("r"))
        rpr = etree.SubElement(run, tag("rPr"))
        fonts = etree.SubElement(rpr, tag("rFonts"))
        for family in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(tag(family), "微软雅黑")
        set_val(rpr, "b", "1" if bold else "0")
        set_val(rpr, "bCs", "1" if bold else "0")
        set_val(rpr, "color", "414141" if kind in {"title", "meta", "tech"} else "000000")
        size = 19 if kind in {"meta", "tech"} else 21
        set_val(rpr, "sz", size)
        set_val(rpr, "szCs", size)
        t = etree.SubElement(run, tag("t"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = content
    return paragraph


with ZipFile(SOURCE) as source_zip:
    document = etree.fromstring(source_zip.read("word/document.xml"))
    table = document.find("w:body/w:tbl", NS)
    rows = table.findall("w:tr", NS)
    original_work_row = rows[6]
    original_cell = original_work_row.find("w:tc", NS)
    original_paragraphs = original_cell.findall("w:p", NS)
    assert len(original_paragraphs) == 23
    assert "基于LangChain" in text_of(original_paragraphs[0])
    assert "智能旅行助手" in text_of(original_paragraphs[6])
    assert "反诈对话" in text_of(original_paragraphs[11])
    assert "数据仓库开发" in text_of(original_paragraphs[17])

    new_content = [
        ("title", "基于多模态大模型与 RAG 的工业图纸智能审查系统"),
        ("meta", "**职责：**整体 AI 应用开发 | 2026.06 - 至今（计划 2026.12 结束）"),
        ("body", "面向弹簧制造场景，开发覆盖图纸识别、标准化审查、人工确认与 CAD 生图对接的 AI 工作台。"),
        ("tech", "**技术栈：**Python / FastAPI / Qwen / RAGFlow / OCR / PostgreSQL / Docker"),
        ("bullet", "接入 **Qwen 多模态模型**提取 PDF/图片中的弹簧尺寸、公差、材料与技术要求，按 JSON 输出约束解析字段，转换为统一候选数据。"),
        ("bullet", "融合 **OCR、几何证据与视觉识别**结果，按来源权重及置信度完成候选选择、冲突检测与证据追溯；支持百度 OCR 异常降级至本地 RapidOCR。"),
        ("bullet", "对接 **RAGFlow**，结合图纸参数和所选标准构造检索问题，通过标准文档过滤与混合检索获取条款，生成带引用依据的标准化建议，支持对话式参数调整。"),
        ("bullet", "结合**工程规则校核与人工确认**，检查尺寸、公差及载荷关系，管理参数联动与确认状态；通过异步 Worker、参数快照和幂等控制，对接 **SolidWorks** 生图及状态回调。"),
    ]
    new_paragraphs = [
        new_paragraph(original_paragraphs[1] if kind == "bullet" else original_paragraphs[0], text, kind=kind, keep_next=index < len(new_content) - 1)
        for index, (kind, text) in enumerate(new_content)
    ]

    def project_row(paragraphs, *, original=False):
        row = deepcopy(original_work_row)
        trpr = row.find(tag("trPr"))
        if trpr is None:
            trpr = etree.Element(tag("trPr"))
            row.insert(0, trpr)
        for h in trpr.findall(tag("trHeight")):
            trpr.remove(h)
        set_val(trpr, "cantSplit", "1")
        cell = row.find(tag("tc"))
        for child in list(cell):
            if child.tag != tag("tcPr"):
                cell.remove(child)
        for index, paragraph in enumerate(paragraphs):
            paragraph = deepcopy(paragraph)
            if original:
                ppr = paragraph.find(tag("pPr"))
                if ppr is None:
                    ppr = etree.Element(tag("pPr"))
                    paragraph.insert(0, ppr)
                set_val(ppr, "keepNext", "0")
                set_val(ppr, "keepLines", "1")
                if index == 0:
                    set_val(ppr, "jc", "left")
            cell.append(paragraph)
        return row

    insertion = table.index(original_work_row)
    replacements = [project_row(new_paragraphs)]
    for start, stop in [(0, 6), (6, 11), (11, 17), (17, 23)]:
        replacements.append(project_row(original_paragraphs[start:stop], original=True))
    table.remove(original_work_row)
    for offset, row in enumerate(replacements):
        table.insert(insertion + offset, row)

    # Keep section labels with the following content, using the retained template.
    for row_index in [2, 5, 7, 9]:
        for paragraph in rows[row_index].xpath(".//w:p", namespaces=NS):
            ppr = paragraph.find(tag("pPr"))
            if ppr is None:
                ppr = etree.Element(tag("pPr"))
                paragraph.insert(0, ppr)
            set_val(ppr, "keepNext", "1")

    updated = etree.tostring(document, encoding="UTF-8", xml_declaration=True, standalone=True)
    with ZipFile(OUTPUT, "w") as output_zip:
        for info in source_zip.infolist():
            output_zip.writestr(copy(info), updated if info.filename == "word/document.xml" else source_zip.read(info.filename))

    with ZipFile(OUTPUT) as final_zip:
        changed = [name for name in source_zip.namelist() if source_zip.read(name) != final_zip.read(name)]
        assert changed == ["word/document.xml"], changed

# All original paragraph texts, media, styles, and related files are retained.
before = Document(SOURCE)
after = Document(OUTPUT)
seen = set()
old_texts = []
for row in before.tables[0].rows:
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        old_texts += [p.text for p in cell.paragraphs if p.text.strip()]
all_text = "\n".join(p.text for row in after.tables[0].rows for cell in row.cells for p in cell.paragraphs)
assert all(text in all_text for text in old_texts)
assert "整体 AI 应用开发" in all_text and "2026.12" in all_text
assert len(after.inline_shapes) == len(before.inline_shapes)
print("OUTPUT", OUTPUT)
print("SOURCE_SHA256", sha256(SOURCE.read_bytes()).hexdigest())
print("ORIGINAL_PARAGRAPHS_PRESERVED", len(old_texts))
print("CHANGED_PARTS", changed)
print("ROWS", len(after.tables[0].rows))
